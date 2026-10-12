#!/usr/bin/env python3
"""Excel de la CNBV -> JSON por mes para el Radar de Sanciones.

Uso:
    python scripts/procesar_excel.py RUTA_DEL_EXCEL.xlsx
    python scripts/procesar_excel.py RUTA_DEL_EXCEL.xlsx --meses 2026-10
    python scripts/procesar_excel.py RUTA_DEL_EXCEL.xlsx --desde 2026-01 --salida public/data

Fuente única: el Excel oficial de la CNBV (hoja «Sanciones CNBV», encabezados en la fila 8).
`mes` = mes de la Fecha de Publicación (AAAA-MM).

Salida (no se editan a mano):
    public/data/sanciones/AAAA-MM.json   registros del mes
    public/data/meses.json               índice de meses (se reconstruye desde los archivos)

Es idempotente: cada mes procesado se reescribe completo, así que correrlo dos veces
no duplica nada. Al final compara registros y suma de montos de cada JSON contra el
Excel y termina con error si algo no cuadra o si aparece un tipo de sanción no previsto.
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

HOJA = "Sanciones CNBV"
FILA_ENCABEZADOS = 8
COLUMNAS = [
    "ID", "Entidad/Persona", "Sector", "Subsector", "Tipo de Sanción",
    "Monto", "Conducta Sancionada", "Fecha de Imposición", "Fecha de Publicación",
]
TIPOS = {
    "Multa (Sanción Pecuniaria)": "Multa",
    "Amonestación": "Amonestación",
}
MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
            "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def limpiar(texto):
    """Colapsa espacios y saltos de línea."""
    return re.sub(r"\s+", " ", str(texto)).strip()


def limpiar_conducta(texto):
    """Además de colapsar espacios, quita el punto final (para agrupar conductas)."""
    return limpiar(texto).rstrip(".").strip()


def parse_monto(valor):
    if isinstance(valor, (int, float)):
        return float(valor)
    return float(str(valor).replace("$", "").replace(",", "").strip())


def parse_fecha(valor):
    if isinstance(valor, datetime):
        return valor
    return datetime.strptime(str(valor).strip(), "%d/%m/%Y")


def num(monto):
    """Entero cuando no hay centavos, para mantener los montos como en la app."""
    return int(monto) if float(monto).is_integer() else round(monto, 2)


def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f"ERROR: no encuentro la hoja «{HOJA}». Hojas: {wb.sheetnames}")
    ws = wb[HOJA]
    filas = ws.iter_rows(min_row=FILA_ENCABEZADOS, values_only=True)
    encabezados = [limpiar(c) if c is not None else "" for c in next(filas)]
    if encabezados[: len(COLUMNAS)] != COLUMNAS:
        sys.exit(f"ERROR: encabezados inesperados en la fila {FILA_ENCABEZADOS}: {encabezados}")

    registros, tipos_raros, errores, ids = [], Counter(), [], Counter()
    for n, f in enumerate(filas, start=FILA_ENCABEZADOS + 1):
        if f[0] in (None, ""):
            continue
        try:
            id_ = int(f[0])
            tipo_xl = limpiar(f[4])
            publicacion = parse_fecha(f[8])
            imposicion = parse_fecha(f[7])
            monto = parse_monto(f[5])
        except Exception as e:  # fila ilegible: se reporta, no se adivina
            errores.append((n, f[:9], str(e)))
            continue
        ids[id_] += 1
        reg = {
            "id": id_,
            "mes": publicacion.strftime("%Y-%m"),
            "entidad": limpiar(f[1]),
            "subsector": limpiar(f[3]),
            "tipo_excel": tipo_xl,
            "monto": num(monto),
            "conducta": limpiar_conducta(f[6]),
            "fechaImposicion": imposicion.strftime("%Y-%m-%d"),
            "_publicacion": publicacion,
        }
        if tipo_xl in TIPOS:
            reg["tipo"] = TIPOS[tipo_xl]
        else:
            tipos_raros[(reg["mes"], tipo_xl)] += 1
        registros.append(reg)
    duplicados = sorted(i for i, c in ids.items() if c > 1)
    return registros, tipos_raros, errores, duplicados


def nombre_mes(mes):
    anio, m = mes.split("-")
    return f"{MESES_ES[int(m) - 1].capitalize()} {anio}"


def fecha_larga(fecha):
    return f"{fecha.day} de {MESES_ES[fecha.month - 1]} de {fecha.year}"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel", help="Ruta del Excel de la CNBV")
    ap.add_argument("--meses", help="Meses a procesar, separados por coma (AAAA-MM). Por defecto, todos desde --desde")
    ap.add_argument("--desde", default="2026-01", help="Primer mes a procesar si no se da --meses (default 2026-01)")
    ap.add_argument("--salida", default=str(Path(__file__).resolve().parent.parent / "public" / "data"),
                    help="Carpeta de salida (default public/data)")
    args = ap.parse_args()

    registros, tipos_raros, errores, duplicados = leer_excel(args.excel)
    print(f"Excel: {len(registros)} registros leídos")

    pedidos = [m.strip() for m in args.meses.split(",")] if args.meses else None
    por_mes = defaultdict(list)
    for r in registros:
        if pedidos is not None:
            if r["mes"] in pedidos:
                por_mes[r["mes"]].append(r)
        elif r["mes"] >= args.desde:
            por_mes[r["mes"]].append(r)
    if pedidos:
        for m in pedidos:
            if m not in por_mes:
                sys.exit(f"ERROR: el Excel no tiene registros publicados en {m}")
    if not por_mes:
        sys.exit("ERROR: no hay meses que procesar")

    problemas = []
    if errores:
        problemas.append(f"{len(errores)} filas ilegibles (primeras: {errores[:3]})")
    if duplicados:
        problemas.append(f"IDs repetidos en el Excel: {duplicados[:10]}")
    raros = {k: v for k, v in tipos_raros.items() if k[0] in por_mes}
    if raros:
        problemas.append("Tipos de sanción no previstos (no se adivinan): " +
                         "; ".join(f"{m} «{t}» x{n}" for (m, t), n in sorted(raros.items())))
    if problemas:
        print("\nPROBLEMAS EN EL EXCEL — no escribo nada:")
        for p in problemas:
            print(" -", p)
        sys.exit(1)

    salida = Path(args.salida)
    carpeta = salida / "sanciones"
    carpeta.mkdir(parents=True, exist_ok=True)

    esperado = {}
    publicaciones = {}
    for mes in sorted(por_mes):
        regs = sorted(por_mes[mes], key=lambda r: r["id"])
        fechas = {r["_publicacion"] for r in regs}
        publicaciones[mes] = max(fechas)
        salida_regs = [{k: r[k] for k in ("id", "mes", "entidad", "subsector", "tipo", "monto", "conducta", "fechaImposicion")}
                       for r in regs]
        with open(carpeta / f"{mes}.json", "w", encoding="utf-8") as fh:
            json.dump(salida_regs, fh, ensure_ascii=False, indent=1)
            fh.write("\n")
        esperado[mes] = (len(regs), sum(r["monto"] for r in regs))

    # Índice: se reconstruye desde TODOS los archivos de la carpeta
    pub_previas = {}
    indice_path = salida / "meses.json"
    if indice_path.exists():
        for m in json.load(open(indice_path, encoding="utf-8")):
            pub_previas[m["id"]] = m.get("fechaPublicacion", "")
    indice = []
    for archivo in sorted(carpeta.glob("*.json")):
        mes = archivo.stem
        regs = json.load(open(archivo, encoding="utf-8"))
        fecha_pub = fecha_larga(publicaciones[mes]) if mes in publicaciones else pub_previas.get(mes, "")
        indice.append({"id": mes, "nombre": nombre_mes(mes), "fechaPublicacion": fecha_pub, "registros": len(regs)})
    with open(indice_path, "w", encoding="utf-8") as fh:
        json.dump(indice, fh, ensure_ascii=False, indent=1)
        fh.write("\n")

    # Validación: lo escrito vs. el Excel
    print(f"\n{'Mes':8} {'Reg. JSON':>9} {'Reg. Excel':>10} {'Suma JSON':>16} {'Suma Excel':>16} {'Multas':>6} {'Amon.':>5} {'Entid.':>6}  OK")
    fallas = 0
    for mes in sorted(por_mes):
        regs = json.load(open(carpeta / f"{mes}.json", encoding="utf-8"))
        n_xl, suma_xl = esperado[mes]
        suma_json = sum(r["monto"] for r in regs)
        ok = len(regs) == n_xl and abs(suma_json - suma_xl) < 0.5
        fallas += 0 if ok else 1
        multas = sum(r["tipo"] == "Multa" for r in regs)
        amon = sum(r["tipo"] == "Amonestación" for r in regs)
        entidades = len({r["entidad"] for r in regs})
        print(f"{mes:8} {len(regs):>9} {n_xl:>10} {suma_json:>16,.0f} {suma_xl:>16,.0f} {multas:>6} {amon:>5} {entidades:>6}  {'sí' if ok else 'NO'}")
    print(f"\nÍndice: {indice_path} ({len(indice)} meses)")
    if fallas:
        sys.exit("ERROR: hay meses que no cuadran contra el Excel")
    print("Todos los meses cuadran contra el Excel.")


if __name__ == "__main__":
    main()
